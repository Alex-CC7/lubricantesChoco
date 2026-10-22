import os
import tempfile
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import letter, landscape
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.units import inch
from datetime import datetime

AZUL = "1A3A5C"
AZUL_CLARO = "D6E4F0"

def estilo_header_excel(ws, columnas, fila=1):
    fill = PatternFill("solid", fgColor=AZUL)
    font = Font(color="FFFFFF", bold=True)
    border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )
    for col_num, titulo in enumerate(columnas, 1):
        cell = ws.cell(row=fila, column=col_num, value=titulo)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = border

def estilo_fila_excel(ws, fila, num_cols, par=True):
    fill = PatternFill("solid", fgColor=AZUL_CLARO if par else "FFFFFF")
    border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )
    for col in range(1, num_cols + 1):
        ws.cell(row=fila, column=col).fill = fill
        ws.cell(row=fila, column=col).border = border
        ws.cell(row=fila, column=col).alignment = Alignment(horizontal='left', vertical='center')

def generar_excel_ventas(ventas):
    wb = Workbook()
    ws = wb.active
    ws.title = "Reporte de Ventas"
    ws.row_dimensions[1].height = 30
    ws.merge_cells('A1:H1')
    titulo = ws['A1']
    titulo.value = f"REPORTE DE VENTAS — {datetime.now().strftime('%d/%m/%Y %H:%M')}"
    titulo.font = Font(bold=True, size=14, color="FFFFFF")
    titulo.fill = PatternFill("solid", fgColor=AZUL)
    titulo.alignment = Alignment(horizontal='center', vertical='center')

    columnas = ['#', 'N° Venta', 'Fecha', 'Cliente', 'Total (Bs)', 'Descuento (Bs)', 'Tipo Venta', 'Vendedor']
    estilo_header_excel(ws, columnas, fila=2)
    ws.column_dimensions['A'].width = 5
    ws.column_dimensions['B'].width = 12
    ws.column_dimensions['C'].width = 20
    ws.column_dimensions['D'].width = 22
    ws.column_dimensions['E'].width = 15
    ws.column_dimensions['F'].width = 15
    ws.column_dimensions['G'].width = 15
    ws.column_dimensions['H'].width = 20

    for i, v in enumerate(ventas, 1):
        fila = i + 2
        estilo_fila_excel(ws, fila, 8, i % 2 == 0)
        ws.cell(row=fila, column=1, value=i)
        ws.cell(row=fila, column=2, value=v['idVenta'])
        ws.cell(row=fila, column=3, value=str(v['FechaVenta']))
        ws.cell(row=fila, column=4, value=v['Cliente'])
        ws.cell(row=fila, column=5, value=float(v['TotalPagar']))
        ws.cell(row=fila, column=6, value=float(v['Descuento']))
        tipo = 'Contado' if str(v['TipoVenta']) == '1' else 'Crédito'
        ws.cell(row=fila, column=7, value=tipo)
        ws.cell(row=fila, column=8, value=v['Usuario'])

    path = os.path.join(tempfile.gettempdir(), 'reporte_ventas.xlsx')
    wb.save(path)
    return path


def generar_pdf_ventas(ventas):
    path = os.path.join(tempfile.gettempdir(), 'reporte_ventas.pdf')
    doc = SimpleDocTemplate(path, pagesize=landscape(letter),
                            rightMargin=30, leftMargin=30, topMargin=40, bottomMargin=30)
    styles = getSampleStyleSheet()
    elements = []

    title_style = ParagraphStyle('title', fontSize=16, fontName='Helvetica-Bold',
                                 textColor=colors.HexColor(f'#{AZUL}'), spaceAfter=10, alignment=1)
    elements.append(Paragraph(f"REPORTE DE VENTAS", title_style))
    elements.append(Paragraph(f"Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')}", styles['Normal']))
    elements.append(Spacer(1, 0.2 * inch))

    data = [['#', 'N° Venta', 'Fecha', 'Cliente', 'Total (Bs)', 'Descuento', 'Tipo', 'Vendedor']]
    for i, v in enumerate(ventas, 1):
        tipo = 'Contado' if str(v['TipoVenta']) == '1' else 'Crédito'
        data.append([
            str(i), str(v['idVenta']), str(v['FechaVenta']),
            v['Cliente'], f"Bs {v['TotalPagar']}", f"Bs {v['Descuento']}", tipo, v['Usuario']
        ])

    t = Table(data, repeatRows=1)
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor(f'#{AZUL}')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 10),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.HexColor(f'#{AZUL_CLARO}'), colors.white]),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('FONTSIZE', (0, 1), (-1, -1), 9),
        ('ROWHEIGHT', (0, 0), (-1, -1), 18),
    ]))
    elements.append(t)
    doc.build(elements)
    return path


def generar_excel_productos(productos):
    wb = Workbook()
    ws = wb.active
    ws.title = "Reporte de Productos"
    ws.row_dimensions[1].height = 30
    ws.merge_cells('A1:G1')
    titulo = ws['A1']
    titulo.value = f"REPORTE DE PRODUCTOS — {datetime.now().strftime('%d/%m/%Y %H:%M')}"
    titulo.font = Font(bold=True, size=14, color="FFFFFF")
    titulo.fill = PatternFill("solid", fgColor=AZUL)
    titulo.alignment = Alignment(horizontal='center', vertical='center')

    columnas = ['#', 'Nombre', 'Marca', 'Stock', 'Precio Venta', 'Proveedor', 'Categoría']
    estilo_header_excel(ws, columnas, fila=2)
    anchos = [5, 25, 18, 14, 15, 25, 18]
    for i, a in enumerate(anchos, 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = a

    for i, p in enumerate(productos, 1):
        fila = i + 2
        estilo_fila_excel(ws, fila, 7, i % 2 == 0)
        ws.cell(row=fila, column=1, value=i)
        ws.cell(row=fila, column=2, value=p['Nombre'])
        ws.cell(row=fila, column=3, value=p['Marca'])
        ws.cell(row=fila, column=4, value=p['Stock'])
        ws.cell(row=fila, column=5, value=float(p['PrecioVenta']))
        ws.cell(row=fila, column=6, value=p['NombreProveedor'])
        ws.cell(row=fila, column=7, value=p['Categoria'])

    path = os.path.join(tempfile.gettempdir(), 'reporte_productos.xlsx')
    wb.save(path)
    return path


def generar_pdf_productos(productos):
    path = os.path.join(tempfile.gettempdir(), 'reporte_productos.pdf')
    doc = SimpleDocTemplate(path, pagesize=landscape(letter),
                            rightMargin=30, leftMargin=30, topMargin=40, bottomMargin=30)
    styles = getSampleStyleSheet()
    elements = []

    title_style = ParagraphStyle('title', fontSize=16, fontName='Helvetica-Bold',
                                 textColor=colors.HexColor(f'#{AZUL}'), spaceAfter=10, alignment=1)
    elements.append(Paragraph("REPORTE DE PRODUCTOS", title_style))
    elements.append(Paragraph(f"Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')}", styles['Normal']))
    elements.append(Spacer(1, 0.2 * inch))

    data = [['#', 'Nombre', 'Marca', 'Stock', 'Precio', 'Proveedor', 'Categoría']]
    for i, p in enumerate(productos, 1):
        data.append([
            str(i), p['Nombre'], p['Marca'], p['Stock'],
            f"Bs {p['PrecioVenta']}", p['NombreProveedor'], p['Categoria']
        ])

    t = Table(data, repeatRows=1)
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor(f'#{AZUL}')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 10),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.HexColor(f'#{AZUL_CLARO}'), colors.white]),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('FONTSIZE', (0, 1), (-1, -1), 9),
        ('ROWHEIGHT', (0, 0), (-1, -1), 18),
    ]))
    elements.append(t)
    doc.build(elements)
    return path


def generar_pdf_venta(venta, detalle):
    """PDF individual de una venta (ticket/comprobante)."""
    path = os.path.join(tempfile.gettempdir(), f"venta_{venta['idVenta']}.pdf")
    doc = SimpleDocTemplate(path, pagesize=letter,
                            rightMargin=50, leftMargin=50, topMargin=50, bottomMargin=40)
    styles = getSampleStyleSheet()
    elements = []

    # Encabezado
    title_style = ParagraphStyle('title', fontSize=20, fontName='Helvetica-Bold',
                                 textColor=colors.HexColor(f'#{AZUL}'), spaceAfter=4, alignment=1)
    sub_style   = ParagraphStyle('sub', fontSize=10, textColor=colors.grey, alignment=1, spaceAfter=2)
    elements.append(Paragraph("BD TALLER", title_style))
    elements.append(Paragraph("Comprobante de Venta", sub_style))
    elements.append(Spacer(1, 0.15 * inch))

    # Línea separadora
    elements.append(Table([['']], colWidths=[6.5 * inch],
                          style=[('LINEBELOW', (0,0), (-1,0), 1.5, colors.HexColor(f'#{AZUL}'))]))
    elements.append(Spacer(1, 0.12 * inch))

    # Info de la venta (2 columnas)
    tipo = 'Contado' if str(venta['TipoVenta']) == '1' else 'Crédito'
    info_data = [
        [Paragraph('<b>N° Venta:</b>', styles['Normal']),
         Paragraph(str(venta['idVenta']), styles['Normal']),
         Paragraph('<b>Fecha:</b>', styles['Normal']),
         Paragraph(str(venta['FechaVenta']), styles['Normal'])],
        [Paragraph('<b>Cliente:</b>', styles['Normal']),
         Paragraph(str(venta['Cliente']), styles['Normal']),
         Paragraph('<b>Vendedor:</b>', styles['Normal']),
         Paragraph(str(venta['Usuario']), styles['Normal'])],
        [Paragraph('<b>Tipo Venta:</b>', styles['Normal']),
         Paragraph(tipo, styles['Normal']), '', ''],
    ]
    info_t = Table(info_data, colWidths=[1.2*inch, 2.05*inch, 1.2*inch, 2.05*inch])
    info_t.setStyle(TableStyle([
        ('FONTSIZE', (0,0), (-1,-1), 9),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('ROWHEIGHT', (0,0), (-1,-1), 18),
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor(f'#{AZUL_CLARO}')),
        ('GRID', (0,0), (-1,-1), 0.3, colors.lightgrey),
        ('LEFTPADDING', (0,0), (-1,-1), 8),
    ]))
    elements.append(info_t)
    elements.append(Spacer(1, 0.2 * inch))

    # Tabla de productos
    det_data = [['#', 'Producto', 'Marca', 'Cant.', 'Precio Unit.', 'Subtotal']]
    subtotal_total = 0.0
    for i, d in enumerate(detalle, 1):
        sub = int(d['cantidad']) * float(d['precio'])
        subtotal_total += sub
        det_data.append([
            str(i), d['Producto'], d['Marca'],
            str(d['cantidad']),
            f"Bs {float(d['precio']):.2f}",
            f"Bs {sub:.2f}"
        ])

    # Filas de totales
    descuento = float(venta['Descuento'])
    total     = float(venta['TotalPagar'])
    det_data.append(['', '', '', '', 'Subtotal:', f"Bs {subtotal_total:.2f}"])
    det_data.append(['', '', '', '', 'Descuento:', f"— Bs {descuento:.2f}"])
    det_data.append(['', '', '', '', 'TOTAL:', f"Bs {total:.2f}"])

    n_det    = len(detalle)
    n_total  = len(det_data)

    det_t = Table(det_data, colWidths=[0.35*inch, 2.0*inch, 1.1*inch, 0.6*inch, 1.3*inch, 1.15*inch],
                  repeatRows=1)
    det_t.setStyle(TableStyle([
        # Encabezado
        ('BACKGROUND',   (0,0),  (-1,0),  colors.HexColor(f'#{AZUL}')),
        ('TEXTCOLOR',    (0,0),  (-1,0),  colors.white),
        ('FONTNAME',     (0,0),  (-1,0),  'Helvetica-Bold'),
        ('FONTSIZE',     (0,0),  (-1,0),  9),
        ('ALIGN',        (0,0),  (-1,0),  'CENTER'),
        # Filas de datos
        ('FONTSIZE',     (0,1),  (-1, n_det), 9),
        ('ROWBACKGROUNDS',(0,1), (-1, n_det),
         [colors.HexColor(f'#{AZUL_CLARO}'), colors.white]),
        ('GRID',         (0,0),  (-1, n_det), 0.4, colors.grey),
        ('ALIGN',        (3,1),  (-1, n_det), 'RIGHT'),
        # Filas de totales
        ('FONTSIZE',     (0, n_det+1), (-1,-1), 9),
        ('ALIGN',        (4, n_det+1), (-1,-1), 'RIGHT'),
        ('FONTNAME',     (4, n_det+1), (-1,-1), 'Helvetica-Bold'),
        ('BACKGROUND',   (0, n_total-1), (-1,-1), colors.HexColor(f'#{AZUL}')),
        ('TEXTCOLOR',    (0, n_total-1), (-1,-1), colors.white),
        ('FONTSIZE',     (0, n_total-1), (-1,-1), 10),
        ('ROWHEIGHT',    (0,0),  (-1,-1), 18),
        ('LEFTPADDING',  (0,0),  (-1,-1), 6),
        ('RIGHTPADDING', (0,0),  (-1,-1), 6),
    ]))
    elements.append(det_t)
    elements.append(Spacer(1, 0.3 * inch))

    # Pie de página
    pie_style = ParagraphStyle('pie', fontSize=8, textColor=colors.grey, alignment=1)
    elements.append(Paragraph(f"Documento generado el {datetime.now().strftime('%d/%m/%Y %H:%M')}", pie_style))
    elements.append(Paragraph("Gracias por su preferencia — BD Taller", pie_style))

    doc.build(elements)
    return path


# ─────────────────────────────────────────────
# HELPERS REUTILIZABLES
# ─────────────────────────────────────────────
def _excel_base(titulo, columnas, merge_hasta):
    wb = Workbook()
    ws = wb.active
    ws.row_dimensions[1].height = 30
    ws.merge_cells(f'A1:{merge_hasta}1')
    t = ws['A1']
    t.value = f"{titulo} — {datetime.now().strftime('%d/%m/%Y %H:%M')}"
    t.font  = Font(bold=True, size=14, color="FFFFFF")
    t.fill  = PatternFill("solid", fgColor=AZUL)
    t.alignment = Alignment(horizontal='center', vertical='center')
    estilo_header_excel(ws, columnas, fila=2)
    return wb, ws

def _pdf_base(path, titulo):
    doc = SimpleDocTemplate(path, pagesize=landscape(letter),
                            rightMargin=30, leftMargin=30, topMargin=40, bottomMargin=30)
    styles = getSampleStyleSheet()
    ts = ParagraphStyle('t', fontSize=15, fontName='Helvetica-Bold',
                        textColor=colors.HexColor(f'#{AZUL}'), spaceAfter=6, alignment=1)
    elems = [
        Paragraph(titulo, ts),
        Paragraph(f"Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')}", styles['Normal']),
        Spacer(1, 0.2*inch),
    ]
    return doc, elems

def _pdf_tabla(data):
    t = Table(data, repeatRows=1)
    t.setStyle(TableStyle([
        ('BACKGROUND',    (0,0), (-1,0),  colors.HexColor(f'#{AZUL}')),
        ('TEXTCOLOR',     (0,0), (-1,0),  colors.white),
        ('FONTNAME',      (0,0), (-1,0),  'Helvetica-Bold'),
        ('FONTSIZE',      (0,0), (-1,0),  9),
        ('ALIGN',         (0,0), (-1,-1), 'CENTER'),
        ('ROWBACKGROUNDS',(0,1), (-1,-1), [colors.HexColor(f'#{AZUL_CLARO}'), colors.white]),
        ('GRID',          (0,0), (-1,-1), 0.4, colors.grey),
        ('FONTSIZE',      (0,1), (-1,-1), 8),
        ('ROWHEIGHT',     (0,0), (-1,-1), 17),
        ('LEFTPADDING',   (0,0), (-1,-1), 5),
    ]))
    return t

def _col_letter(n):
    return chr(64 + n)


# ─────────────────────────────────────────────
# CATEGORÍAS
# ─────────────────────────────────────────────
def generar_excel_categorias(categorias):
    cols = ['#', 'Nombre', 'Descripción', 'Últ. Actualización']
    wb, ws = _excel_base("REPORTE DE CATEGORÍAS", cols, _col_letter(len(cols)))
    for w, c in zip([5,25,35,25], ['A','B','C','D']):
        ws.column_dimensions[c].width = w
    for i, cat in enumerate(categorias, 1):
        r = i + 2
        estilo_fila_excel(ws, r, 4, i%2==0)
        ws.cell(r,1,i); ws.cell(r,2,cat['Nombre'])
        ws.cell(r,3,cat['Descripcion']); ws.cell(r,4,str(cat['FechaActualizacion']))
    path = os.path.join(tempfile.gettempdir(), 'categorias.xlsx')
    wb.save(path); return path

def generar_pdf_categorias(categorias):
    path = os.path.join(tempfile.gettempdir(), 'categorias.pdf')
    doc, elems = _pdf_base(path, "REPORTE DE CATEGORÍAS")
    data = [['#','Nombre','Descripción','Últ. Actualización']]
    for i, c in enumerate(categorias,1):
        data.append([str(i), c['Nombre'], c['Descripcion'], str(c['FechaActualizacion'])])
    elems.append(_pdf_tabla(data)); doc.build(elems); return path


# ─────────────────────────────────────────────
# CLIENTES
# ─────────────────────────────────────────────
def generar_excel_clientes(clientes):
    cols = ['#','Razón Social','NIT/CI','Últ. Actualización']
    wb, ws = _excel_base("REPORTE DE CLIENTES", cols, _col_letter(len(cols)))
    for w, c in zip([5,35,20,25], ['A','B','C','D']):
        ws.column_dimensions[c].width = w
    for i, cl in enumerate(clientes,1):
        r = i+2; estilo_fila_excel(ws,r,4,i%2==0)
        ws.cell(r,1,i); ws.cell(r,2,cl['RazonSocial'])
        ws.cell(r,3,cl['Nit/CI']); ws.cell(r,4,str(cl['FechaActualizacion']))
    path = os.path.join(tempfile.gettempdir(),'clientes.xlsx')
    wb.save(path); return path

def generar_pdf_clientes(clientes):
    path = os.path.join(tempfile.gettempdir(),'clientes.pdf')
    doc, elems = _pdf_base(path,"REPORTE DE CLIENTES")
    data = [['#','Razón Social','NIT/CI','Últ. Actualización']]
    for i,c in enumerate(clientes,1):
        data.append([str(i),c['RazonSocial'],str(c['Nit/CI']),str(c['FechaActualizacion'])])
    elems.append(_pdf_tabla(data)); doc.build(elems); return path


# ─────────────────────────────────────────────
# USUARIOS
# ─────────────────────────────────────────────
def generar_excel_usuarios(usuarios):
    cols = ['#','Nombre Completo','CI','Celular','Usuario','Rol']
    wb, ws = _excel_base("REPORTE DE USUARIOS", cols, _col_letter(len(cols)))
    for w,c in zip([5,35,15,15,20,15],['A','B','C','D','E','F']):
        ws.column_dimensions[c].width = w
    for i,u in enumerate(usuarios,1):
        r=i+2; estilo_fila_excel(ws,r,6,i%2==0)
        nombre = f"{u['Nombre']} {u['PrimerApellido']} {u['SegundoApellido']}"
        rol = 'Admin' if u['Rol']=='1' else 'Personal'
        ws.cell(r,1,i); ws.cell(r,2,nombre); ws.cell(r,3,u['CI'])
        ws.cell(r,4,u['NroCelular']); ws.cell(r,5,u['NombreUsuario']); ws.cell(r,6,rol)
    path = os.path.join(tempfile.gettempdir(),'usuarios.xlsx')
    wb.save(path); return path

def generar_pdf_usuarios(usuarios):
    path = os.path.join(tempfile.gettempdir(),'usuarios.pdf')
    doc, elems = _pdf_base(path,"REPORTE DE USUARIOS")
    data = [['#','Nombre Completo','CI','Celular','Usuario','Rol']]
    for i,u in enumerate(usuarios,1):
        nombre = f"{u['Nombre']} {u['PrimerApellido']} {u['SegundoApellido']}"
        data.append([str(i),nombre,str(u['CI']),str(u['NroCelular']),
                     u['NombreUsuario'],'Admin' if u['Rol']=='1' else 'Personal'])
    elems.append(_pdf_tabla(data)); doc.build(elems); return path
